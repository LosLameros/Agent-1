"""Sestaví excel s modelem DIP (vystupy/grafy/2026-10-21-dip-model.xlsx) – všechny výpočty jako vzorce.

Spouštěj z kořene repozitáře: python3 vystupy/grafy/2026-10-21-dip-model-sestaveni.py
Potom excel přepočítej (LibreOffice / recalc.py) a spusť vystupy/grafy/2026-10-21-dip-grafy.py.
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

VYSTUP = "vystupy/grafy/2026-10-21-dip-model.xlsx"
wb = Workbook()


def F(**k):
    return Font(name="Arial", size=k.pop("size", 10), **k)


MODRA, CERNA, ZELENA, TUCNE = F(color="0000FF"), F(), F(color="008000"), F(bold=True)
NADPIS = F(size=14, bold=True)
POZN = F(italic=True, color="666666")
ZLUTA = PatternFill("solid", fgColor="FFFF00")
HLAV = PatternFill("solid", fgColor="ECEBF8")
KC = '#,##0" Kč";-#,##0" Kč";"-"'
PCT = "0.0%"
PCT2 = "0.00%"
PCT4 = "0.0000%"


def hlavicka(ws, radek, nazvy, sirky):
    for j, (h, w) in enumerate(zip(nazvy, sirky), 1):
        c = ws.cell(radek, j, h)
        c.font, c.fill = TUCNE, HLAV
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[radek].height = 44


# --- Předpoklady ---------------------------------------------------------------
ws = wb.active
ws.title = "Předpoklady"
ws["A1"] = "Model DIP – předpoklady"
ws["A1"].font = NADPIS
ws["A2"] = "Modré buňky se žlutým podkladem jsou vstupy – po změně se celý sešit přepočítá."
ws["A2"].font = F(italic=True)
VSTUPY = [
    (4, "Měsíční vklad", 4000, KC, True, "Zadání (Ája B., Marek P.): 4 000 Kč měsíčně, vklad vždy na začátku měsíce."),
    (5, "Očekávaný roční výnos (před poplatky Portu)", 0.07, PCT, True, "Zadání: průměrný růst 7 % ročně, stejný pro DIP i investici mimo DIP."),
    (6, "Roční limit odpočtu příspěvků na DIP", 48000, KC, True, "Zákon o daních z příjmů, § 15a: společný limit 48 000 Kč ročně (DIP, penzijní spoření, životní pojištění)."),
    (7, "Sazba daně z příjmů", 0.15, PCT, True, "Zadání: daňová úspora 7 200 Kč ročně = 48 000 Kč × 15 %."),
    (8, "Měsíc vrácení daně v následujícím roce", 4, "0", True, "Přeplatek na dani se vrací typicky v dubnu po podání přiznání / ročním zúčtování."),
    (9, "Poplatek DIP (ročně, účtovaný měsíčně)", 0.003, PCT2, True, "Zadání: DIP má poplatek 0,3 % ročně, účtuje se měsíčně (1/12 z hodnoty majetku)."),
    (10, "Sleva na poplatku za fixaci (mimo DIP)", 0.40, PCT, True, "Zadání: investice mimo DIP s fixací se slevou 40 % (Program pro dlouhodobé investory; 40 % odpovídá fixaci na 15 let). Platí pro horizonty 20 a 30 let."),
    (11, "Sleva na poplatku za fixaci – 10letý horizont", 0.30, PCT, True, "Zadání: u 10letého horizontu je poplatek mimo DIP s fixací 0,7 % ročně = 1 % × (1 − 30 %); s objemem klesá stejně jako u 40% slevy."),
    (12, "Měsíční výnos (odvozený)", "=(1+B5)^(1/12)-1", PCT4, False, "Měsíční výnos tak, aby za 12 měsíců dal přesně roční výnos z B5."),
    (13, "Roční vklady", "=B4*12", KC, False, ""),
    (14, "Roční daňová úspora", "=MIN(B13,B6)*B7", KC, False, "Odečitatelné příspěvky (max. limit) × sazba daně."),
]
for r, popis, hodnota, fmt, vstup, pozn in VSTUPY:
    ws.cell(r, 1, popis).font = CERNA
    c = ws.cell(r, 2, hodnota)
    c.number_format = fmt
    c.font = MODRA if vstup else CERNA
    if vstup:
        c.fill = ZLUTA
    ws.cell(r, 3, pozn).font = POZN

ws["A16"] = "Poplatek za správu mimo DIP (Portfolia od Portu)"
ws["A16"].font = TUCNE
for j, h in enumerate(["Majetek od (Kč)", "Sazba ročně", "Po slevě (20 a 30 let)", "Po slevě (10 let)"], 1):
    c = ws.cell(17, j, h)
    c.font, c.fill = TUCNE, HLAV
PASMA = [(0, 0.01), (500000, 0.008), (1000000, 0.006), (5000000, 0.004)]
for i, (od, sazba) in enumerate(PASMA):
    r = 18 + i
    a = ws.cell(r, 1, od)
    a.number_format, a.font, a.fill = KC, MODRA, ZLUTA
    b = ws.cell(r, 2, sazba)
    b.number_format, b.font, b.fill = PCT2, MODRA, ZLUTA
    c = ws.cell(r, 3, f"=B{r}*(1-$B$10)")
    c.number_format = PCT2
    c = ws.cell(r, 4, f"=B{r}*(1-$B$11)")
    c.number_format = PCT2
ws["A22"] = ("Zdroj: Portu magazín, Nový sazebník Portu (23. 4. 2024, platí od 1. 5. 2024), "
             "https://magazin.portu.cz/novy-sazebnik-portu-je-jednodussi-a-umoznuje-investovat-levneji/ – "
             "sazba pásma platí na celý majetek; poplatek účtovaný měsíčně (1/12 z hodnoty majetku).")
ws["A22"].font = POZN

ws["A24"] = "Scénáře"
ws["A24"].font = TUCNE
SCEN = [
    ("Bez DIP s fixací", "Portfolio od Portu s fixací mimo DIP: stejný vklad a výnos, poplatek podle pásem se slevou za fixaci, bez daňové úspory. Výnos z cenných papírů držených déle než 3 roky je osvobozen od daně, takže se nedaní ani mimo DIP."),
    ("DIP + daňová úspora", "Investice v DIP s poplatkem 0,3 % + vrácená daň za každý rok ponechaná jako hotovost (bez dalšího zhodnocení)."),
    ("DIP + reinvestovaná daňová úspora", "Investice v DIP + vrácená daň vložená do DIP v měsíci vrácení (stejný výnos i poplatek 0,3 %). Další odpočet nepřináší – limit 48 000 Kč vyčerpají samotné vklady."),
]
for i, (n, p) in enumerate(SCEN):
    ws.cell(25 + i, 1, n).font = CERNA
    ws.cell(25 + i, 3, p).font = POZN
ws["A29"] = "Další předpoklady"
ws["A29"].font = TUCNE
for i, t in enumerate([
    "Vklad na začátku měsíce; zhodnocení i poplatek měsíčně (poplatek = 1/12 roční sazby z hodnoty majetku na konci měsíce).",
    "Daňová úspora za poslední rok horizontu se vrací až v dubnu následujícího roku – na konci horizontu je započtená jako nárok.",
    "Podmínky DIP: výběr nejdříve po 120 měsících a zároveň po 60. roce věku. Horizonty 10/20/30 let odpovídají investorům ve věku 50/40/30 let.",
    "Bez inflace a bez příspěvků zaměstnavatele. Nákladovost podkladových ETF je součástí výnosu 7 %.",
]):
    ws.cell(30 + i, 1, "• " + t).font = F(color="333333")
ws.column_dimensions["A"].width = 46
ws.column_dimensions["B"].width = 16
ws.column_dimensions["C"].width = 120

P = "'Předpoklady'!"
# --- Měsíce --------------------------------------------------------------------
wm = wb.create_sheet("Měsíce")
hlavicka(wm, 1, ["Měsíc", "Rok", "Vklad",
                 "Bez DIP: hodnota před poplatkem", "Bez DIP: sazba poplatku ročně", "Bez DIP s fixací: hodnota (konec měsíce)",
                 "DIP: hodnota investice (konec měsíce)", "Vrácená daň v měsíci", "Vrácená daň kumulovaně (hotovost)",
                 "Reinvestovaná daňová úspora v DIP (hodnota)", "Poplatky bez DIP kumulovaně", "Poplatky DIP kumulovaně",
                 "Bez DIP (10 let): hodnota před poplatkem", "Bez DIP (10 let): sazba poplatku ročně",
                 "Bez DIP s fixací (10 let): hodnota (konec měsíce)", "Poplatky bez DIP (10 let) kumulovaně"],
         [7, 5, 11, 15, 11, 16, 16, 12, 15, 17, 14, 14, 15, 11, 16, 14])
for m in range(1, 361):
    r = m + 1
    pr = r - 1
    wm.cell(r, 1, m)
    wm.cell(r, 2, f"=ROUNDUP(A{r}/12,0)")
    wm.cell(r, 3, f"={P}$B$4").font = ZELENA
    wm.cell(r, 4, f"=({'F' + str(pr) if m > 1 else '0'}+C{r})*(1+{P}$B$12)")
    wm.cell(r, 5, f"=INDEX({P}$C$18:$C$21,MATCH(D{r},{P}$A$18:$A$21,1))")
    wm.cell(r, 6, f"=D{r}*(1-E{r}/12)")
    wm.cell(r, 7, f"=({'G' + str(pr) if m > 1 else '0'}+C{r})*(1+{P}$B$12)*(1-{P}$B$9/12)")
    wm.cell(r, 8, f"=IF(AND(A{r}>12,MOD(A{r},12)={P}$B$8),{P}$B$14,0)")
    wm.cell(r, 9, f"={'I' + str(pr) if m > 1 else '0'}+H{r}")
    wm.cell(r, 10, f"={'J' + str(pr) if m > 1 else '0'}*(1+{P}$B$12)*(1-{P}$B$9/12)+H{r}")
    wm.cell(r, 11, f"={'K' + str(pr) if m > 1 else '0'}+D{r}*E{r}/12")
    wm.cell(r, 12, f"={'L' + str(pr) if m > 1 else '0'}+(({'G' + str(pr) if m > 1 else '0'}+C{r})"
                   f"+{'J' + str(pr) if m > 1 else '0'})*(1+{P}$B$12)*{P}$B$9/12")
    # stejný výpočet bez DIPu se slevou pro 10letý horizont (sazby ze sloupce D tabulky poplatků)
    wm.cell(r, 13, f"=({'O' + str(pr) if m > 1 else '0'}+C{r})*(1+{P}$B$12)")
    wm.cell(r, 14, f"=INDEX({P}$D$18:$D$21,MATCH(M{r},{P}$A$18:$A$21,1))")
    wm.cell(r, 15, f"=M{r}*(1-N{r}/12)")
    wm.cell(r, 16, f"={'P' + str(pr) if m > 1 else '0'}+M{r}*N{r}/12")
    for j in (3, 4, 6, 7, 8, 9, 10, 11, 12, 13, 15, 16):
        wm.cell(r, j).number_format = KC
    for j in (5, 14):
        wm.cell(r, j).number_format = PCT2
wm.freeze_panes = "A2"

M = "'Měsíce'!"
# --- Roky (podklad grafů) ------------------------------------------------------
wr = wb.create_sheet("Roky")
hlavicka(wr, 1, ["Rok", "Měsíc", "Vloženo celkem", "Bez DIP s fixací", "DIP + daňová úspora",
                 "DIP + reinvestovaná daňová úspora", "Výhoda DIP oproti fixaci", "Výhoda DIP s reinvestováním",
                 "Poplatky bez DIP kumulovaně", "Poplatky DIP kumulovaně",
                 "Bez DIP s fixací (sleva pro 10 let)", "Výhoda DIP oproti fixaci (10 let)",
                 "Výhoda DIP s reinvestováním (10 let)", "Poplatky bez DIP kumulovaně (10 let)"],
         [6, 8, 16, 18, 18, 22, 18, 20, 16, 16, 18, 18, 20, 16])
for t in range(0, 31):
    r = t + 2
    wr.cell(r, 1, t)
    wr.cell(r, 2, f"=A{r}*12")
    wr.cell(r, 3, f"=B{r}*{P}$B$4")
    wr.cell(r, 4, f"=IF(A{r}=0,0,INDEX({M}$F$2:$F$361,B{r}))")
    wr.cell(r, 5, f"=IF(A{r}=0,0,INDEX({M}$G$2:$G$361,B{r})+INDEX({M}$I$2:$I$361,B{r})+{P}$B$14)")
    wr.cell(r, 6, f"=IF(A{r}=0,0,INDEX({M}$G$2:$G$361,B{r})+INDEX({M}$J$2:$J$361,B{r})+{P}$B$14)")
    wr.cell(r, 7, f"=E{r}-D{r}")
    wr.cell(r, 8, f"=F{r}-D{r}")
    wr.cell(r, 9, f"=IF(A{r}=0,0,INDEX({M}$K$2:$K$361,B{r}))")
    wr.cell(r, 10, f"=IF(A{r}=0,0,INDEX({M}$L$2:$L$361,B{r}))")
    wr.cell(r, 11, f"=IF(A{r}=0,0,INDEX({M}$O$2:$O$361,B{r}))")
    wr.cell(r, 12, f"=E{r}-K{r}")
    wr.cell(r, 13, f"=F{r}-K{r}")
    wr.cell(r, 14, f"=IF(A{r}=0,0,INDEX({M}$P$2:$P$361,B{r}))")
    for j in range(3, 15):
        wr.cell(r, j).number_format = KC
wr["A34"] = "Hodnoty ke konci roku. Daňová úspora za právě skončený rok je započtená jako nárok (vrací se v dubnu dalšího roku)."
wr["A34"].font = POZN
wr.freeze_panes = "A2"

# --- Výsledky ------------------------------------------------------------------
wv = wb.create_sheet("Výsledky", 1)
wv["A1"] = "DIP – výsledky pro horizonty 10, 20 a 30 let"
wv["A1"].font = NADPIS
hlavicka(wv, 3, ["Horizont (let)", "Vloženo celkem", "Bez DIP s fixací", "DIP + daňová úspora",
                 "DIP + reinvestovaná daňová úspora", "Výhoda DIP oproti fixaci", "Výhoda DIP s reinvestováním",
                 "Poplatky bez DIP", "Poplatky DIP", "Pro investora ve věku"],
         [14, 16, 18, 20, 24, 20, 22, 16, 16, 16])
R = "'Roky'!"
for i, n in enumerate((10, 20, 30)):
    r = 4 + i
    wv.cell(r, 1, n).font = MODRA
    # 10letý horizont: bez DIPu se slevou pro 10 let (sloupce K–N listu Roky), 20 a 30 let se slevou 40 %
    for j, sl in enumerate("CKEFLMNJ" if n == 10 else "CDEFGHIJ", 2):
        c = wv.cell(r, j, f"=INDEX({R}${sl}$2:${sl}$32,A{r}+1)")
        c.number_format, c.font = KC, ZELENA
    wv.cell(r, 10, f'=60-A{r}&" let"')
wv["A8"] = "Výpočet: list Měsíce (měsíční model) → Roky (stav ke konci roku) → Výsledky. Předpoklady a zdroje: list Předpoklady."
wv["A8"].font = POZN

for s in wb.worksheets:
    for row in s.iter_rows():
        for c in row:
            if c.value is not None and c.font.name != "Arial":
                c.font = F(size=c.font.size or 10, bold=c.font.bold, italic=c.font.italic, color=c.font.color)
wb.save(VYSTUP)
print("uloženo", VYSTUP)
