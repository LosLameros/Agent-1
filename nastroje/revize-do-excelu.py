#!/usr/bin/env python3
# Použití: python3 nastroje/revize-do-excelu.py <vstup.json> <vystup.xlsx>
# Sestaví Excel revize klientských portfolií ve formátu vzoru
# examples/revize-portfolii/Revize_portfolia_vzor.xlsx (listy Přehled, Detail, Překryvy, Kurzy).
# Všechny CZK hodnoty a procenta jsou vzorce (odkazy na list Kurzy a na CELKEM v Přehledu),
# takže změna kurzu v listu Kurzy přepočítá celý sešit. Skript zároveň vypíše spočtené
# hodnoty pro textovou část revize. Struktura vstupního JSON je popsaná
# v .claude/agents/revize-portfolia.md (kapitola 5).
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ZELENA = "FF00A03C"     # záhlaví a titulek
MANTINEL = "FFD9EAD3"   # mantinely (nedotýkat se)
PREKRYV = "FFFFF3CD"    # překryvy / duplicitní expozice
CELKEM = "FFE9E2F3"     # součtové řádky
ZDANITELNE = "FFF4CCCC" # prodej s daňovou povinností / překročený limit
LIMIT_CP = 100000       # hodnotový test: roční limit hrubých příjmů z prodeje CP (Kč)
SAZBA = 0.15            # základní sazba daně z příjmů FO (orientační výpočet)
NEUVEDENO = "neuvedeno"

TENKA = Side(style="thin", color="FFBFBFBF")
RAMECEK = Border(top=TENKA, bottom=TENKA, left=TENKA, right=TENKA)


def vypln(barva):
    return PatternFill("solid", fgColor=barva)


def font(**kw):
    return Font(name="Arial", size=kw.pop("size", 10), **kw)


def titulek(ws, text, sloupcu):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=sloupcu)
    c = ws.cell(1, 1, text)
    c.font = font(size=14, bold=True, color="FFFFFFFF")
    c.fill = vypln(ZELENA)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24


def podtitulek(ws, text, sloupcu, vyska=30):
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=sloupcu)
    c = ws.cell(2, 1, text)
    c.font = font(size=9, color="FF333333")
    c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = vyska


def zahlavi(ws, radek, nazvy):
    for i, n in enumerate(nazvy, 1):
        c = ws.cell(radek, i, n)
        c.font = font(bold=True, color="FFFFFFFF")
        c.fill = vypln(ZELENA)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = RAMECEK


def bunka(ws, r, s, hodnota, fmt=None, zarovnani="left", barva=None, tucne=False):
    c = ws.cell(r, s, hodnota)
    c.font = font(bold=tucne)
    c.border = RAMECEK
    c.alignment = Alignment(horizontal=zarovnani, vertical="center", wrap_text=zarovnani == "left")
    if fmt:
        c.number_format = fmt
    if barva:
        c.fill = vypln(barva)
    return c


def poznamky(ws, od_radku, texty, sloupcu):
    r = od_radku
    for t in texty:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=sloupcu)
        c = ws.cell(r, 1, t)
        c.font = font(size=9, color="FF333333")
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.row_dimensions[r].height = max(15, 13 * (len(t) // 120 + 1))
        r += 1
    return r


def sirky(ws, hodnoty):
    for pismeno, w in zip("ABCDEFG", hodnoty):
        ws.column_dimensions[pismeno].width = w
    # tisk na šířku, všechny sloupce na jednu stránku
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def main(vstup, vystup):
    d = json.load(open(vstup, encoding="utf-8"))
    kurzy = {k.upper(): float(v) for k, v in d["kurzy"]["hodnoty"].items()}
    kurzy["CZK"] = 1.0

    def czk(mena, hodnota):
        if hodnota is None:
            return None
        mena = mena.upper()
        if mena not in kurzy:
            sys.exit(f"Chybí kurz pro měnu {mena} (doplň ho do kurzy.hodnoty).")
        return hodnota * kurzy[mena]

    wb = Workbook()
    prehled = wb.active
    prehled.title = "Přehled"
    detail = wb.create_sheet("Detail")
    prekryvy = wb.create_sheet("Překryvy")
    list_kurzy = wb.create_sheet("Kurzy")

    # ---------- Kurzy ----------
    k = d["kurzy"]
    list_kurzy["A1"] = "Přepočet měn – použité kurzy"
    list_kurzy["A1"].font = font(size=12, bold=True)
    radky_kurzu = {}
    r = 3
    list_kurzy.cell(r, 1, "Datum").font = font(bold=True)
    list_kurzy.cell(r, 2, k["datum"]).font = font()
    r += 1
    for mena in [m for m in ["USD", "EUR"] if m in kurzy] + sorted(m for m in kurzy if m not in ("USD", "EUR", "CZK")):
        list_kurzy.cell(r, 1, f"1 {mena} =").font = font(bold=True)
        c = list_kurzy.cell(r, 2, kurzy[mena])
        c.font = font()
        c.number_format = '#,##0.000 "Kč"'
        c.alignment = Alignment(horizontal="left")
        radky_kurzu[mena] = r
        r += 1
    list_kurzy.cell(r, 1, "Měna přepočtu").font = font(bold=True)
    list_kurzy.cell(r, 2, "CZK").font = font()
    r += 1
    list_kurzy.cell(r, 1, "Zdroj").font = font(bold=True)
    c = list_kurzy.cell(r, 2, k["zdroj"])
    c.font = font()
    c.alignment = Alignment(wrap_text=True, vertical="top")
    r += 2
    for t in d.get("poznamky_kurzy", []):
        c = list_kurzy.cell(r, 1, t)
        c.font = font(size=9, color="FF333333")
        r += 1
    sirky(list_kurzy, [16, 70])

    def vzorec_czk(mena, adresa_hodnoty):
        mena = mena.upper()
        if mena == "CZK":
            return f"={adresa_hodnoty}"
        return f"={adresa_hodnoty}*Kurzy!$B${radky_kurzu[mena]}"

    # ---------- Přehled ----------
    pozice = d["pozice"]
    for p in pozice:
        p["_czk"] = czk(p["mena"], p.get("hodnota"))
    pozice.sort(key=lambda p: (p["_czk"] is None, -(p["_czk"] or 0)))
    celkem = sum(p["_czk"] or 0 for p in pozice)

    titulek(prehled, f"Portu – přehled všech portfolií a přímých pozic (přepočet k {k['datum']})", 6)
    meny = ", ".join(f"1 {m} = {kurzy[m]} Kč" for m in radky_kurzu)
    podtitulek(prehled, f"Kurzy: {meny} ({k.get('zdroj_kratce', k['zdroj'])}). Setříděno od nejvyšší hodnoty. "
                        "Desetinná tečka; % na 1 des. místo.", 6)
    zahlavi(prehled, 4, ["Portfolio", "Instrument (název + ticker/ISIN)", "Původní měna",
                         "Hodnota (pův. měna)", "Hodnota v CZK", "Podíl na celku"])
    prvni = 5
    radek_celkem = prvni + len(pozice)
    adresa_pozice = {}
    for i, p in enumerate(pozice):
        r = prvni + i
        barva = MANTINEL if p.get("mantinel") else (PREKRYV if p.get("prekryv") else None)
        bunka(prehled, r, 1, p["portfolio"], barva=barva)
        bunka(prehled, r, 2, p["instrument"], barva=barva)
        bunka(prehled, r, 3, p["mena"].upper(), zarovnani="center", barva=barva)
        if p.get("hodnota") is None:
            for s in (4, 5, 6):
                bunka(prehled, r, s, NEUVEDENO, zarovnani="right", barva=barva)
        else:
            bunka(prehled, r, 4, p["hodnota"], "#,##0.00", "right", barva)
            bunka(prehled, r, 5, vzorec_czk(p["mena"], f"D{r}"), "#,##0.00", "right", barva)
            bunka(prehled, r, 6, f"=E{r}/$E${radek_celkem}", "0.0%", "right", barva)
        if p.get("id"):
            adresa_pozice[p["id"]] = f"'Přehled'!E{r}"
    for s in range(1, 7):
        bunka(prehled, radek_celkem, s, None, barva=CELKEM, tucne=True)
    prehled.cell(radek_celkem, 1, "CELKEM")
    prehled.cell(radek_celkem, 5, f"=SUM(E{prvni}:E{radek_celkem - 1})").number_format = "#,##0.00"
    prehled.cell(radek_celkem, 6, f"=SUM(F{prvni}:F{radek_celkem - 1})").number_format = "0.0%"
    for s in (5, 6):
        prehled.cell(radek_celkem, s).alignment = Alignment(horizontal="right")
    poz = ["Zeleně = MANTINELY (nedotýkat se; do součtu a % se ale počítají)."]
    if any(p.get("hodnota") is None for p in pozice):
        poz.append("Pozor: některé hodnoty nejsou na printscreenech čitelné (neuvedeno) – CELKEM je proto neúplný.")
    poznamky(prehled, radek_celkem + 2, poz + d.get("poznamky_prehled", []), 6)
    sirky(prehled, [26, 42, 12, 20, 18, 14])
    prehled.freeze_panes = "A5"
    CELEK = f"'Přehled'!$E${radek_celkem}"

    # ---------- Detail ----------
    poradi = {}
    for p in pozice:
        poradi.setdefault(p["portfolio"], len(poradi))
    slozky = d.get("slozky", [])
    for s in slozky:
        s["_czk"] = czk(s["mena"], s.get("hodnota"))
    slozky.sort(key=lambda s: (poradi.get(s["portfolio"], 999), s["portfolio"],
                               s["_czk"] is None, -(s["_czk"] or 0)))

    titulek(detail, "Detail viditelných složek portfolií", 7)
    podtitulek(detail, d.get("upozorneni_detail",
                             "POZOR: zobrazeny jen složky viditelné na printscreenech. Hodnoty složek jsou "
                             "vnitřní podíly portfolií – NEsčítají se do celku znovu, celkovou hodnotu drží list Přehled."),
               7, vyska=42)
    zahlavi(detail, 4, ["Portfolio", "Instrument / složka", "Původní měna", "Hodnota (pův. měna)",
                        "Hodnota v CZK", "Zastoupení v portfoliu", "Podíl na celku"])
    adresa_slozky = {}
    for i, s in enumerate(slozky):
        r = 5 + i
        barva = PREKRYV if s.get("prekryv") else None
        bunka(detail, r, 1, s["portfolio"], barva=barva)
        bunka(detail, r, 2, s["instrument"], barva=barva)
        bunka(detail, r, 3, s["mena"].upper(), zarovnani="center", barva=barva)
        if s.get("hodnota") is None:
            bunka(detail, r, 4, NEUVEDENO, zarovnani="right", barva=barva)
            bunka(detail, r, 5, NEUVEDENO, zarovnani="right", barva=barva)
            bunka(detail, r, 7, NEUVEDENO, zarovnani="right", barva=barva)
        else:
            bunka(detail, r, 4, s["hodnota"], "#,##0.00", "right", barva)
            bunka(detail, r, 5, vzorec_czk(s["mena"], f"D{r}"), "#,##0.00", "right", barva)
            bunka(detail, r, 7, f"=E{r}/{CELEK}", "0.0%", "right", barva)
        z = s.get("zastoupeni")
        bunka(detail, r, 6, NEUVEDENO if z is None else z, None if z is None else "0.0%", "right", barva)
        if s.get("id"):
            adresa_slozky[s["id"]] = f"Detail!E{r}"
    poznamky(detail, 5 + len(slozky) + 1,
             ["Žlutě = složka, která se objevuje i v jiném portfoliu / jako přímá pozice (překryv – viz list Překryvy)."]
             + d.get("poznamky_detail", []), 7)
    sirky(detail, [22, 40, 12, 20, 18, 20, 14])
    detail.freeze_panes = "A5"

    # ---------- Překryvy ----------
    adresy = {**adresa_pozice, **adresa_slozky}
    hodnoty = {p["id"]: p["_czk"] for p in pozice if p.get("id")}
    hodnoty.update({s["id"]: s["_czk"] for s in slozky if s.get("id")})

    def soucet(refs):
        chybi = [x for x in refs if x not in adresy]
        if chybi:
            sys.exit(f"Překryvy odkazují na neznámá id: {', '.join(chybi)}")
        platne = [x for x in refs if hodnoty[x] is not None]
        vzorec = "=" + "+".join(adresy[x] for x in platne) if platne else NEUVEDENO
        return vzorec, sum(hodnoty[x] for x in platne), len(platne) < len(refs)

    titulek(prekryvy, "Překryvy a duplicitní expozice napříč portfolii", 4)
    podtitulek(prekryvy, "Stejný / velmi podobný instrument sečtený napříč portfolii. CZK přes kurzy z listu Kurzy. "
                         "% z celku (list Přehled).", 4)
    zahlavi(prekryvy, 4, ["Instrument / expozice", "Ve kterých portfoliích (počet)",
                          "Hodnota v CZK (součet)", "Podíl na celku"])
    vypis = []
    radky = []
    for p in d.get("prekryvy", []):
        f, v, neuplne = soucet(p["refs"])
        radky.append((v, p, f, neuplne))
    radky.sort(key=lambda x: -x[0])
    r = 5
    for v, p, f, neuplne in radky:
        bunka(prekryvy, r, 1, p["expozice"], barva=PREKRYV)
        bunka(prekryvy, r, 2, p["portfolia"])
        bunka(prekryvy, r, 3, f, "#,##0.00", "right")
        bunka(prekryvy, r, 4, f"=C{r}/{CELEK}" if f != NEUVEDENO else NEUVEDENO, "0.0%", "right")
        vypis.append((p["expozice"], v, neuplne))
        r += 1
    for blok in d.get("tematicke_bloky", []):
        r += 1
        prekryvy.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        c = bunka(prekryvy, r, 1, blok["nazev"].upper(), barva=PREKRYV, tucne=True)
        r += 1
        od = r
        polozky = []
        for pol in blok["polozky"]:
            f, v, neuplne = soucet(pol["refs"])
            polozky.append((v, pol, f))
        polozky.sort(key=lambda x: -x[0])
        for v, pol, f in polozky:
            bunka(prekryvy, r, 1, "   " + pol["popis"])
            bunka(prekryvy, r, 2, pol["portfolia"])
            bunka(prekryvy, r, 3, f, "#,##0.00", "right")
            bunka(prekryvy, r, 4, f"=C{r}/{CELEK}" if f != NEUVEDENO else NEUVEDENO, "0.0%", "right")
            r += 1
        for s in range(1, 5):
            bunka(prekryvy, r, s, None, barva=CELKEM, tucne=True)
        prekryvy.cell(r, 1, f"{blok['nazev_celkem']}")
        prekryvy.cell(r, 3, f"=SUM(C{od}:C{r - 1})").number_format = "#,##0.00"
        prekryvy.cell(r, 4, f"=C{r}/{CELEK}").number_format = "0.0%"
        for s in (3, 4):
            prekryvy.cell(r, s).alignment = Alignment(horizontal="right")
        vypis.append((blok["nazev_celkem"], sum(x[0] for x in polozky), False))
        r += 1
    poznamky(prekryvy, r + 1, d.get("poznamky_prekryvy", []), 4)
    sirky(prekryvy, [46, 52, 24, 14])
    prekryvy.freeze_panes = "A5"

    # ---------- Daně (volitelné: jen když je k dispozici daňový přehled) ----------
    dane_vypis = None
    if d.get("dane"):
        dn = d["dane"]
        ws = wb.create_sheet("Daně", index=3)
        titulek(ws, "Daně: časový a hodnotový test navržených prodejů (FIFO napříč portfolii)", 7)
        podtitulek(ws, dn.get("upozorneni",
                              "Osvobozená hodnota je z daňového přehledu klienta; neosvobozená část je dopočtena "
                              "(drženo celkem − osvobozeno). Portu prodává metodou FIFO napříč všemi portfolii: "
                              "prodej čerpá nejdřív nejstarší (osvobozené) kusy daného instrumentu, ať je prodej "
                              "z kteréhokoli portfolia. Hodnoty jsou v aktuálních cenách – orientační."), 7, vyska=42)
        zahlavi(ws, 4, ["Instrument (název + ticker/ISIN)", "Ve kterých portfoliích", "Drženo celkem v CZK",
                        "Časově osvobozeno (CZK)", "Neosvobozeno (CZK, dopočteno)", "Osvobozený podíl"])
        bunka_osvob = {}
        r = 5
        for t in dn["instrumenty"]:
            f, v, neuplne = soucet(t["refs"])
            t["_drzeno"] = v
            t["_osvob"] = czk(t.get("mena", "CZK"), t.get("osvobozeno"))
            bunka(ws, r, 1, t["instrument"])
            bunka(ws, r, 2, t.get("portfolia", ""))
            bunka(ws, r, 3, f, "#,##0.00", "right")
            if t["_osvob"] is None:
                for sl in (4, 5, 6):
                    bunka(ws, r, sl, NEUVEDENO, zarovnani="right")
            else:
                hod = t["osvobozeno"]
                vz = hod if t.get("mena", "CZK").upper() == "CZK" else f"={hod}*Kurzy!$B${radky_kurzu[t['mena'].upper()]}"
                bunka(ws, r, 4, vz, "#,##0.00", "right")
                bunka(ws, r, 5, f"=MAX(0,C{r}-D{r})", "#,##0.00", "right")
                bunka(ws, r, 6, f"=IF(C{r}>0,MIN(1,D{r}/C{r}),0)", "0.0%", "right")
            bunka_osvob[t["id"]] = (r, t)
            r += 1

        # navržené prodeje sloučené po instrumentech (FIFO čerpá společnou osvobozenou zásobu)
        prodeje = {}
        for pr in dn.get("prodeje", []):
            if pr["instrument"] not in bunka_osvob:
                sys.exit(f"Prodej odkazuje na neznámý instrument daňového přehledu: {pr['instrument']}")
            x = prodeje.setdefault(pr["instrument"], {"refs": [], "castka": 0.0, "navrhy": [], "zisk": 0.0})
            x["refs"] += pr.get("refs", [])
            # orientační zisk/ztráta části bez časového testu (CZK); chybí-li u kteréhokoli prodeje, je neznámý
            x["zisk"] = None if x["zisk"] is None or pr.get("zisk") is None else x["zisk"] + pr["zisk"]
            x["castka"] += pr.get("castka", 0) or 0
            x["navrhy"].append(pr["navrh"])
        r += 1
        zahlavi(ws, r, ["Prodávaný instrument", "Návrhy", "Příjem z prodeje (CZK)",
                        "Z toho časově osvobozeno (FIFO)", "Z toho bez časového testu", "Daňový stav",
                        "Orientační zisk/ztráta (bez čas. testu)"])
        r += 1
        od = r
        r_suma = od + len(prodeje)          # úhrn navržených prodejů
        r_dalsi, r_uhrn, r_test, r_zdan = r_suma + 1, r_suma + 2, r_suma + 3, r_suma + 4
        r_rz, r_dan_pred, r_dan_po, r_dan_navic = r_suma + 5, r_suma + 6, r_suma + 7, r_suma + 8
        radky_prodeju = []
        for tid, x in prodeje.items():
            radek_t, t = bunka_osvob[tid]
            casti = []
            hodnota = x["castka"]
            if x["refs"]:
                f, v, _ = soucet(x["refs"])
                casti.append(f[1:])
                hodnota += v
            if x["castka"]:
                casti.append(str(x["castka"]))
            if hodnota > t["_drzeno"] + 0.5:
                sys.exit(f"Prodej {t['instrument']} ({hodnota:,.2f} CZK) převyšuje drženou hodnotu "
                         f"({t['_drzeno']:,.2f} CZK) – zkontroluj refs/castka.")
            osv = None if t["_osvob"] is None else min(hodnota, t["_osvob"])
            radky_prodeju.append((radek_t, t, x, casti, hodnota, osv))

        dalsi = dn.get("dalsi_prodeje_v_roce")   # známé další prodeje CP v roce (CZK), None = neznámé
        uhrn = sum(h for *_, h, _ in radky_prodeju) + (dalsi or 0)
        hodnotovy = uhrn <= LIMIT_CP
        nezname = any(o is None for *_, o in radky_prodeju)
        osvobozeno_ct = sum(o for *_, o in radky_prodeju if o is not None)
        zdanitelne = 0.0 if hodnotovy else sum(h - o for *_, h, o in radky_prodeju if o is not None)

        dane_vypis = []
        for radek_t, t, x, casti, hodnota, osv in radky_prodeju:
            if hodnotovy:
                stav = "osvobozeno (hodnotový test)"
            elif osv is None:
                stav = "nelze posoudit – osvobození neuvedeno"
            elif hodnota - osv < 0.005:
                stav = "osvobozeno (časový test)"
            elif osv < 0.005:
                stav = "zdanitelné"
            else:
                stav = "částečně osvobozeno (časový test)"
            barva = None if stav.startswith("osvobozeno") else ZDANITELNE
            bunka(ws, r, 1, t["instrument"], barva=barva)
            bunka(ws, r, 2, ", ".join(x["navrhy"]), barva=barva)
            bunka(ws, r, 3, "=" + "+".join(casti), "#,##0.00", "right", barva)
            test = f"$C${r_uhrn}<={LIMIT_CP}"
            if osv is None:
                bunka(ws, r, 4, NEUVEDENO, zarovnani="right", barva=barva)
                bunka(ws, r, 5, NEUVEDENO, zarovnani="right", barva=barva)
                bunka(ws, r, 6, f'=IF({test},"osvobozeno (hodnotový test)",'
                                f'"nelze posoudit – osvobození neuvedeno")', barva=barva)
            else:
                bunka(ws, r, 4, f"=MIN(C{r},D{radek_t})", "#,##0.00", "right", barva)
                bunka(ws, r, 5, f"=C{r}-D{r}", "#,##0.00", "right", barva)
                bunka(ws, r, 6, f'=IF({test},"osvobozeno (hodnotový test)",IF(E{r}<0.005,'
                                f'"osvobozeno (časový test)",IF(D{r}<0.005,"zdanitelné",'
                                f'"částečně osvobozeno (časový test)")))', barva=barva)
            z = x["zisk"]
            bunka(ws, r, 7, NEUVEDENO if z is None else z, None if z is None else "#,##0.00", "right", barva)
            dane_vypis.append((t["instrument"], ", ".join(x["navrhy"]), hodnota, osv, stav, z))
            r += 1

        def souhrn(radek, popis, barva, hodnoty_sloupcu):
            for sl in range(1, 8):
                bunka(ws, radek, sl, None, barva=barva, tucne=True)
            ws.cell(radek, 1, popis)
            for sl, (hod, fmt) in hodnoty_sloupcu.items():
                c = ws.cell(radek, sl, hod)
                if fmt:
                    c.number_format = fmt
                c.alignment = Alignment(horizontal="right" if fmt else "left", wrap_text=not fmt)

        souhrn(r_suma, "Úhrn navržených prodejů", CELKEM,
               {s: (f"=SUM({pis}{od}:{pis}{r_suma - 1})", "#,##0.00") for s, pis in ((3, "C"), (4, "D"), (5, "E"), (7, "G"))})
        souhrn(r_dalsi, "Další prodeje CP v témže roce (i mimo Portu)", CELKEM,
               {3: (dalsi if dalsi is not None else NEUVEDENO, "#,##0.00"),
                6: ("zadáno" if dalsi is not None else "neznámé – počítáno s nulou", None)})
        souhrn(r_uhrn, "Úhrn příjmů z prodeje CP za rok (hrubě, vč. časově osvobozených)", CELKEM,
               {3: (f"=C{r_suma}+N(C{r_dalsi})", "#,##0.00")})
        souhrn(r_test, f"Hodnotový test {LIMIT_CP:,} Kč".replace(",", " "),
               CELKEM if hodnotovy else ZDANITELNE,
               {6: (f'=IF(C{r_uhrn}<={LIMIT_CP},"SPLNĚN – vše osvobozeno","NESPLNĚN")', None)})
        souhrn(r_zdan, "Zdanitelný příjem (bez časového testu)", ZDANITELNE if zdanitelne > 0 or (nezname and not hodnotovy) else CELKEM,
               {5: (f"=IF(C{r_uhrn}<={LIMIT_CP},0,E{r_suma})", "#,##0.00"),
                6: ("+ část neuvedeno" if nezname and not hodnotovy else ("povinnost podat DP" if zdanitelne > 0 else ""), None)})
        # orientační daň: základ = zisky − ztráty z prodejů CP bez časového testu za rok (ztráty se započítávají)
        rz = dn.get("realizovany_zisk_v_roce")   # dosud realizovaný zisk v roce (bez časového testu), None = neznámý
        zisky = [z for *_, z in dane_vypis]
        zisk_navrhu = None if any(z is None for z in zisky) else sum(zisky)
        hodnotovy_pred = (dalsi or 0) <= LIMIT_CP
        dan_pred = 0.0 if hodnotovy_pred else SAZBA * max(0.0, rz or 0.0)
        dan_po = None if zisk_navrhu is None else (0.0 if hodnotovy else SAZBA * max(0.0, (rz or 0.0) + zisk_navrhu))
        souhrn(r_rz, "Realizovaný zisk z dalších prodejů v roce (bez časového testu)", CELKEM,
               {7: (rz if rz is not None else NEUVEDENO, "#,##0.00"),
                6: ("zadáno" if rz is not None else "neznámý – počítáno s nulou", None)})
        souhrn(r_dan_pred, f"Orientační daň z prodejů CP bez návrhů ({SAZBA:.0%})", CELKEM,
               {7: (f"=IF(N(C{r_dalsi})<={LIMIT_CP},0,{SAZBA}*MAX(0,N(G{r_rz})))", "#,##0.00")})
        souhrn(r_dan_po, f"Orientační daň z prodejů CP s návrhy ({SAZBA:.0%})", CELKEM,
               {7: (f"=IF(C{r_uhrn}<={LIMIT_CP},0,{SAZBA}*MAX(0,N(G{r_rz})+G{r_suma}))"
                    if zisk_navrhu is not None else NEUVEDENO, "#,##0.00")})
        navic = None if dan_po is None else dan_po - dan_pred
        souhrn(r_dan_navic, "DAŇ NAVÍC Z NÁVRHŮ (orientačně)",
               ZDANITELNE if navic is None or navic > 0.5 else CELKEM,
               {7: (f"=G{r_dan_po}-G{r_dan_pred}" if navic is not None else NEUVEDENO, "#,##0.00"),
                6: ("zisk některého prodeje neuveden" if navic is None else "", None)})
        r = r_dan_navic + 2
        pozn = ["Červeně = prodej se zdanitelnou částí nebo s neznámým osvobozením; nesplněný hodnotový test.",
                "Hodnotový test (§ 4 odst. 1 písm. t) ZDP): když úhrn hrubých příjmů z prodeje CP za rok – ze všech "
                "prodejů, i mimo Portu a i časově osvobozených – nepřesáhne 100 000 Kč, je osvobozeno vše. Při "
                "překročení se zdaňují všechny příjmy, které nesplňují časový test (ne jen část nad limit).",
                "Časový test (§ 4 odst. 1 písm. u) ZDP): 3 roky mezi nákupem a prodejem, pro každý kus zvlášť; "
                "Portu páruje metodou FIFO. Zdanitelný příjem ≠ daň: základem je příjem − pořizovací cena − poplatky.",
                f"Orientační daň: {SAZBA:.0%} ze součtu zisků a ztrát z prodejů CP bez časového testu za rok (ztráty "
                "se proti ziskům ve stejném roce započítávají). Zisk je orientační (ze screenů výnosu) – skutečný zisk "
                "určí párování FIFO, případně může jít o sazbu 23 % (základ nad 36násobek průměrné mzdy)."]
        if osvobozeno_ct > 5_000_000:
            pozn.append("POZOR: osvobozené příjmy přesahují 5 mil. Kč za rok – je třeba je oznámit finančnímu úřadu "
                        "do konce lhůty pro podání daňového přiznání.")
        poznamky(ws, r, pozn + dn.get("poznamky", []), 7)
        sirky(ws, [40, 30, 22, 22, 24, 30, 22])
        ws.freeze_panes = "A5"
        dane_vypis = (dn["instrumenty"], dane_vypis, dict(
            uhrn=uhrn, dalsi=dalsi, hodnotovy=hodnotovy, zdanitelne=zdanitelne,
            nezname=nezname, osvobozeno_ct=osvobozeno_ct, rz=rz, zisk_navrhu=zisk_navrhu,
            dan_pred=dan_pred, dan_po=dan_po, navic=navic))

    wb.save(vystup)

    # ---------- Kontrolní výpis pro textovou část revize ----------
    print(f"Uloženo: {vystup}")
    print(f"CELKEM: {celkem:,.2f} CZK")
    print("\nPřehled (od největší):")
    for p in pozice:
        v = p["_czk"]
        print(f"  {p['portfolio']} | {p['instrument']} | "
              + (NEUVEDENO if v is None else f"{v:,.2f} CZK | {v / celkem:.1%}")
              + (" | MANTINEL" if p.get("mantinel") else ""))
    print("\nPřekryvy:")
    for nazev, v, neuplne in vypis:
        print(f"  {nazev} | {v:,.2f} CZK | {v / celkem:.1%}" + (" | NEÚPLNÉ (část neuvedeno)" if neuplne else ""))
    soucty = {}
    for s in slozky:
        soucty[s["portfolio"]] = soucty.get(s["portfolio"], 0) + (s["_czk"] or 0)
    print("\nKontrola: viditelné složky vs. hodnota portfolia v Přehledu:")
    for p in pozice:
        if p["portfolio"] in soucty and p["_czk"]:
            print(f"  {p['portfolio']}: složky {soucty[p['portfolio']]:,.2f} / portfolio {p['_czk']:,.2f} CZK "
                  f"({soucty[p['portfolio']] / p['_czk']:.1%})")

    if dane_vypis:
        instrumenty, prodeje, sh = dane_vypis
        print("\nČasový test (drženo / osvobozeno / neosvobozeno):")
        for t in instrumenty:
            o = t["_osvob"]
            print(f"  {t['instrument']}: {t['_drzeno']:,.2f} / "
                  + (NEUVEDENO if o is None else f"{o:,.2f} / {max(0, t['_drzeno'] - o):,.2f} CZK"))
        print("\nNavržené prodeje (příjem / časově osvobozeno / bez časového testu):")
        for nazev, navrhy, h, o, stav, z in prodeje:
            print(f"  {nazev} [{navrhy}]: {h:,.2f} / "
                  + (f"{o:,.2f} / {h - o:,.2f} CZK" if o is not None else NEUVEDENO) + f" | {stav}"
                  + f" | orient. zisk {NEUVEDENO if z is None else f'{z:,.2f} CZK'}")
        print(f"  Úhrn příjmů z prodeje CP za rok: {sh['uhrn']:,.2f} CZK"
              + (" (další prodeje v roce neznámé – počítáno s nulou)" if sh["dalsi"] is None else ""))
        print(f"  Hodnotový test {LIMIT_CP:,} Kč: " + ("SPLNĚN – vše osvobozeno" if sh["hodnotovy"] else "NESPLNĚN"))
        print(f"  Zdanitelný příjem: {sh['zdanitelne']:,.2f} CZK"
              + (" + část neuvedeno" if sh["nezname"] and not sh["hodnotovy"] else "")
              + (" → povinnost podat daňové přiznání" if sh["zdanitelne"] > 0 else ""))
        print(f"  Realizovaný zisk v roce dosud: " + (NEUVEDENO if sh["rz"] is None else f"{sh['rz']:,.2f} CZK")
              + " | zisk/ztráta z návrhů: " + (NEUVEDENO if sh["zisk_navrhu"] is None else f"{sh['zisk_navrhu']:,.2f} CZK"))
        print(f"  Orientační daň ({SAZBA:.0%}) bez návrhů: {sh['dan_pred']:,.2f} CZK | s návrhy: "
              + (NEUVEDENO if sh["dan_po"] is None else f"{sh['dan_po']:,.2f} CZK")
              + " | DAŇ NAVÍC Z NÁVRHŮ: " + (NEUVEDENO if sh["navic"] is None else f"{sh['navic']:,.2f} CZK"))
        if sh["osvobozeno_ct"] > 5_000_000:
            print("  POZOR: osvobozené příjmy > 5 mil. Kč → oznamovací povinnost vůči FÚ")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit("Použití: python3 nastroje/revize-do-excelu.py <vstup.json> <vystup.xlsx>")
    main(sys.argv[1], sys.argv[2])
